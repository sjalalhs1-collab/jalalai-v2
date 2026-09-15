import sys,json,os
from pathlib import Path

def main():
    spec=Path(sys.argv[1]) if len(sys.argv)>1 else None
    r=json.loads(spec.read_text()) if spec else json.load(sys.stdin); out=Path(r['out']); out.parent.mkdir(parents=True,exist_ok=True); kind=r['kind']; content=r.get('content',''); title=r.get('title') or out.stem
    if kind=='docx':
        from docx import Document
        d=Document(); d.add_heading(title,0)
        for line in content.splitlines(): d.add_paragraph(line)
        d.save(out)
    elif kind=='xlsx':
        from openpyxl import Workbook
        wb=Workbook(); ws=wb.active; ws.title=(r.get('sheet') or 'Data')[:31]
        for row in content.strip().splitlines(): ws.append(row.split(','))
        wb.save(out)
    elif kind=='pptx':
        from pptx import Presentation
        prs=Presentation(); slides=r.get('slides') or [content]
        for i,text in enumerate(slides):
            layout=prs.slide_layouts[0 if i==0 else 1]; sl=prs.slides.add_slide(layout)
            if sl.shapes.title: sl.shapes.title.text=title if i==0 else f'{title} — {i+1}'
            if len(sl.placeholders)>1: sl.placeholders[1].text=text
        prs.save(out)
    else: raise SystemExit('unsupported')
    print(json.dumps({'path':str(out),'bytes':out.stat().st_size}))
if __name__=='__main__': main()
